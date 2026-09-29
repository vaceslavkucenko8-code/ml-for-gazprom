"""Read the source workbook without changing it; preserve raw and derived data."""
import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from pathlib import Path
from urllib.parse import urlsplit

import openpyxl

HEADERS = ['№', 'Технология (слабый сигнал)', 'Область', 'Компании',
           'Почему это слабый сигнал', 'Стадия развития', 'Тренд упоминаний',
           'Балл (стадия+тренд)', 'Источники']
FIELDS = ['signal_id', 'technology', 'domain', 'companies_raw', 'rationale_raw',
          'stage_raw', 'trend_raw', 'reference_score', 'sources_raw']
STAGE_PATTERNS = {
    'concept': r'концепц', 'research': r'исследован',
    'prototype': r'прототип|\bpoc\b', 'pilot': r'пилот',
    'early_adoption': r'ранн\w*\s+(?:внедр|постав|сери)|первые\s+(?:внедрен|интеграц|постав)',
    'mass_adoption': r'массов\w*\s+(?:внедрен|производств)|широк\w*\s+внедрен',
}


def normalize(value):
    return re.sub(r'\s+', ' ', str(value)).strip()


def stage_features(text):
    """Return mentioned stages, not a claim about actual market maturity."""
    text = normalize(text).casefold()
    flags = {f'stage_{key}': bool(re.search(pattern, text))
             for key, pattern in STAGE_PATTERNS.items()}
    flags['stage_multiple'] = sum(flags.values()) > 1
    flags['stage_unknown'] = not any(flags.values())
    flags['stage_has_transition'] = '→' in text
    return flags


def trend_category(text):
    prefix = re.split(r'[:—–]', normalize(text).casefold(), maxsplit=1)[0].strip()
    mapping = {'растёт быстро': 'fast_growth', 'растёт': 'growth',
               'стабильный/растёт': 'stable_or_growth',
               'растёт (медленно)': 'slow_growth',
               'стабильный с ускорением': 'stable_accelerating'}
    return mapping.get(prefix, 'unknown')


def source_links(text):
    # The source uses Markdown links. URLs have no literal parentheses here.
    return [{'title_raw': title, 'url': url, 'host': urlsplit(url).netloc.casefold()}
            for title, url in re.findall(r'\[([^\]]+)\]\((https?://[^\s)]+)\)', text)]


def read_workbook(path):
    book = openpyxl.load_workbook(path, data_only=False)
    try:
        if 'Слабые сигналы' not in book.sheetnames:
            raise ValueError('Required sheet is missing: Слабые сигналы')
        sheet = book['Слабые сигналы']
        headers = [sheet.cell(2, c).value for c in range(2, 11)]
        if headers != HEADERS:
            raise ValueError(f'Unexpected headers in B2:J2: {headers!r}')
        rows = []
        for row in sheet.iter_rows(min_row=3, min_col=2, max_col=10):
            values = [cell.value for cell in row]
            if all(value is None for value in values):
                continue
            if any(cell.data_type == 'f' for cell in row):
                raise ValueError(f'Unexpected formula at Excel row {row[0].row}')
            record = dict(zip(FIELDS, values))
            record['excel_row'] = row[0].row
            rows.append(record)
        metadata = {'sheets': book.sheetnames, 'sheet': sheet.title,
                    'dimensions': sheet.calculate_dimension(),
                    'merged_ranges': [str(r) for r in sheet.merged_cells.ranges]}
        return rows, metadata
    finally:
        book.close()


def validate(rows):
    if not rows:
        raise ValueError('No data records found')
    for record in rows:
        for name in FIELDS:
            value = record[name]
            if value is None or isinstance(value, str) and not value.strip():
                raise ValueError(f'Missing {name} at row {record["excel_row"]}')
            if name in ('signal_id', 'reference_score'):
                if type(value) is not int:
                    raise ValueError(f'{name} must be an integer at row {record["excel_row"]}')
            elif not isinstance(value, str):
                raise ValueError(f'{name} must be text at row {record["excel_row"]}')
    ids = [record['signal_id'] for record in rows]
    if len(ids) != len(set(ids)):
        raise ValueError('Duplicate signal_id')


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def prepare(input_path, output_dir):
    checksum = hashlib.sha256(input_path.read_bytes()).hexdigest()
    raw, metadata = read_workbook(input_path)
    validate(raw)
    clean, features, links = [], [], []
    for record in raw:
        item = {key: normalize(value) if isinstance(value, str) else value
                for key, value in record.items()}
        sources = source_links(item['sources_raw'])
        raw_url_count = len(re.findall(r'https?://[^\s)]+', item['sources_raw']))
        if len(sources) != raw_url_count or not sources:
            raise ValueError(f'Unparsed source links for signal {item["signal_id"]}')
        item['sources'] = sources
        clean.append(item)
        features.append({'signal_id': item['signal_id'], 'technology': item['technology'],
                         'domain': item['domain'], **stage_features(item['stage_raw']),
                         'trend_category': trend_category(item['trend_raw'])})
        links.extend({'signal_id': item['signal_id'], **link} for link in sources)
    urls = Counter(link['url'] for link in links)
    titles = Counter(item['technology'].casefold() for item in clean)
    scores = [item['reference_score'] for item in clean]
    report = {**metadata, 'records': len(raw), 'source_fields': len(FIELDS),
              'missing_by_field': {field: 0 for field in FIELDS},
              'types_by_field': {field: dict(Counter(type(r[field]).__name__ for r in raw)) for field in FIELDS},
              'duplicate_rows_excluding_id': len(raw) - len({tuple(r[f] for f in FIELDS if f != 'signal_id') for r in raw}),
              'duplicate_normalized_titles': {k: v for k, v in titles.items() if v > 1},
              'domain_counts': dict(Counter(r['domain'] for r in clean)),
              'score_counts': dict(Counter(scores)),
              'scores_sorted_descending': scores == sorted(scores, reverse=True),
              'stage_raw_counts': dict(Counter(r['stage_raw'] for r in clean)),
              'trend_counts': dict(Counter(r['trend_category'] for r in features)),
              'stage_review_ids': [r['signal_id'] for r in features if r['stage_unknown'] or r['stage_multiple']],
              'source_url_occurrences': len(links), 'unique_urls': len(urls),
              'unique_hosts': len({link['host'] for link in links}),
              'repeated_urls': {url: count for url, count in urls.items() if count > 1},
              'warnings': ['No binary ground-truth column in supplied workbook; hidden labels are expected by the specification.',
                           'Stage flags describe text mentions, not verified maturity.',
                           'Trend categories are not measured publication time series.',
                           'URLs extracted, not fetched or independently verified.',
                           'signal_id is a join key, never a model input.',
                           'Exclude reference_score and rationale_raw from detector inputs; evaluate stage/trend leakage separately.']}
    output_dir.mkdir(parents=True, exist_ok=True)
    write_json(output_dir / 'signals_raw.json', raw)
    write_json(output_dir / 'signals_clean.json', clean)
    write_json(output_dir / 'sources.json', links)
    write_json(output_dir / 'quality_report.json', report)
    write_json(output_dir / 'manifest.json', {'schema_version': '1.0', 'feature_version': '0.1',
               'input_file': input_path.name, 'input_sha256': checksum,
               'openpyxl_version': openpyxl.__version__, 'records': len(raw),
               'model_input_columns': [key for key in features[0] if key != 'signal_id'],
               'excluded_fields': ['signal_id', 'excel_row', 'reference_score', 'rationale_raw', 'companies_raw', 'sources_raw'],
               'status': 'data preparation only; no trained model or validated detection metrics'})
    with (output_dir / 'features.csv').open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=list(features[0]))
        writer.writeheader()
        writer.writerows(features)
    if hashlib.sha256(input_path.read_bytes()).hexdigest() != checksum:
        raise RuntimeError('Source workbook changed during processing')
    return report


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--input', type=Path, required=True)
    parser.add_argument('--output', type=Path, default=Path('data/processed'))
    args = parser.parse_args()
    report = prepare(args.input, args.output)
    print(f'Records: {report["records"]}; URL occurrences: {report["source_url_occurrences"]}; unique URLs: {report["unique_urls"]}')


if __name__ == '__main__':
    main()
